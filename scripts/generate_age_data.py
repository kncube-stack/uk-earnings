from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


PERIOD_FILES = {
    "annual": "PROV - Age Group Table 6.7a   Annual pay - Gross 2026.xlsx",
    "weekly": "PROV - Age Group Table 6.1a   Weekly pay - Gross 2026.xlsx",
    "hourly": "PROV - Age Group Table 6.5a   Hourly Pay - Gross 2026.xlsx",
    "hours": "PROV - Age Group Table 6.9a   Paid hours worked - Total 2026.xlsx",
}

SHEET_MAP = {
    "All": "all",
    "Male": "male",
    "Female": "female",
    "Full-Time": "ft",
    "Male Full-Time": "male_ft",
    "Female Full-Time": "female_ft",
}

AGE_LABELS = {
    "16-17b": "16–17",
    "18-21": "18–21",
    "22-29": "22–29",
    "30-39": "30–39",
    "40-49": "40–49",
    "50-59": "50–59",
    "60+": "60+",
}

FIELDS = [
    ("median", 4),
    ("mean", 6),
    ("p10", 8),
    ("p20", 9),
    ("p25", 10),
    ("p30", 11),
    ("p40", 12),
    ("p60", 13),
    ("p70", 14),
    ("p75", 15),
    ("p80", 16),
    ("p90", 17),
]


def clean_value(value):
    if value in (None, "", "x", "..", ":", "-"):
        return None
    if isinstance(value, str):
        stripped = value.strip().lstrip("\'").lstrip("'")
        if stripped in {"", "x", "..", ":", "-"}:
            return None
        try:
            numeric = float(stripped)
        except ValueError:
            return stripped
        if numeric.is_integer():
            return int(numeric)
        return round(numeric, 2)
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def parse_sheet(path: Path, sheet_name: str):
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[sheet_name]
    rows = []

    for row in sheet.iter_rows(min_row=6, values_only=True):
        source_label = clean_value(row[0])
        if source_label not in AGE_LABELS:
            continue
        rows.append(
            {
                "label": AGE_LABELS[source_label],
                **{field: clean_value(row[index - 1]) for field, index in FIELDS},
            }
        )

    return rows


def build_data(source_dir: Path):
    data = {}
    for period, filename in PERIOD_FILES.items():
        workbook_path = source_dir / filename
        if not workbook_path.exists():
            raise FileNotFoundError(workbook_path)
        data[period] = {
            key: parse_sheet(workbook_path, sheet_name)
            for sheet_name, key in SHEET_MAP.items()
        }
    return data


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage: generate_age_data.py /path/to/table6/folder")
    data = build_data(Path(sys.argv[1]))
    print("// Generated from ONS ASHE 2026 Table 6 workbooks.")
    print("// Do not edit by hand; run scripts/generate_age_data.py instead.\n")
    print(f"export const AGE_DATA = {json.dumps(data, ensure_ascii=False, separators=(',', ':'))};")


if __name__ == "__main__":
    main()
