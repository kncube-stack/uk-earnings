from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


FILES = {
    "gross": "PROV - Age Group Table 6.1a   Weekly pay - Gross 2026.xlsx",
    "basic": "PROV - Age Group Table 6.3a   Basic Pay - Including other pay 2026.xlsx",
    "overtime": "PROV - Age Group Table 6.4a   Overtime pay 2026.xlsx",
    "incentive": "PROV - Age Group Table 6.8a   Annual pay - Incentive 2026.xlsx",
}

SHEET_MAP = {
    "All": "all",
    "Male": "male",
    "Female": "female",
    "Full-Time": "ft",
    "Male Full-Time": "male_ft",
    "Female Full-Time": "female_ft",
}

AGE_LABELS = ["16-17b", "18-21", "22-29", "30-39", "40-49", "50-59", "60+"]


def clean_number(value):
    if value in (None, "", "x", "..", ":", "-"):
        return None
    if isinstance(value, str):
        value = value.strip().lstrip("'")
        if value in {"", "x", "..", ":", "-"}:
            return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else round(number, 2)


def means_by_age(path: Path, sheet_name: str):
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[sheet_name]
    result = {}
    for row in sheet.iter_rows(min_row=6, values_only=True):
        label = row[0].strip().lstrip("'") if isinstance(row[0], str) else row[0]
        if label in AGE_LABELS:
            result[label] = clean_number(row[5])
    return result


def build_data(source_dir: Path):
    paths = {name: source_dir / filename for name, filename in FILES.items()}
    missing = [str(path) for path in paths.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing workbook(s): " + ", ".join(missing))

    data = {}
    for sheet_name, key in SHEET_MAP.items():
        series = {name: means_by_age(path, sheet_name) for name, path in paths.items()}
        rows = []
        for age in AGE_LABELS:
            gross = series["gross"].get(age)
            basic = series["basic"].get(age)
            overtime = series["overtime"].get(age)
            other = None
            if gross is not None and basic is not None:
                other = round(gross - basic - (overtime or 0), 1)
            rows.append(
                {
                    "basic": basic,
                    "ot": overtime,
                    "other": other,
                    "inc": series["incentive"].get(age),
                }
            )
        data[key] = rows
    return data


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage: generate_age_composition_data.py /path/to/table6/folder")
    data = build_data(Path(sys.argv[1]))
    print("// Generated from ONS ASHE 2026 Table 6 workbooks.")
    print("// Do not edit by hand; run scripts/generate_age_composition_data.py instead.\n")
    print(f"export const AGE_COMPOSITION = {json.dumps(data, ensure_ascii=False, separators=(',', ':'))};")


if __name__ == "__main__":
    main()
